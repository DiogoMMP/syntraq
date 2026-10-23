"""Render a study document from markdown to .docx (pandoc) and .pdf (Word / LibreOffice).

Usage:
    python render.py <source.md> [--no-pdf] [--reference <reference.docx>]

The .docx and .pdf are written next to the source, with the same name. The source's
frontmatter may set:
    title, subtitle, date, lang, toc-title   used by pandoc for the cover and the table of contents
    toc: false                               leave the table of contents out (e.g. an exam paper)
    version, updated                         shown in the footer of every page, when present
    section-breaks: false                    don't start every top-level (#) section on a new page

Layout: every top-level section starts on a new page; tables (centred, full grid, coloured header, alternating row colours), figures and captions are centred; paragraphs, code blocks and the boxes are never
cut across pages; a table that fits on a page is kept whole (a longer one breaks only between rows and
repeats its header row). "Fits" means at most 12 rows and about 1200 characters.

Wiki links ([[Article]] / [[Article|text]]) are turned into plain text, since they mean
nothing outside Obsidian. The band and state emoji (🔥 ⚠️ 📎 🔴 🟡 🟢 ⚪) are swapped for plain symbols
in the Word/PDF only: a single emoji makes Word embed the whole emoji font (~4 MB per PDF).

Diagrams: a fenced block ```mermaid (optionally followed by a caption on the same line, e.g.
```mermaid Bayes network) is rendered to a PNG with mermaid-cli (`mmdc`) and put in the document as a
figure. Images (`![caption](figuras/x.png){width=80%}`) are embedded as they are.

Exit codes: 0 ok (the PDF may still be skipped — see the last line printed), 2 bad input,
3 pandoc not installed, 4 pandoc failed, 5 the source has mermaid diagrams and mmdc is not installed,
6 mmdc could not render a diagram.
"""

import argparse
import json
import os
import platform
import re
import shutil
import struct
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import make_reference

WIKILINK = re.compile(r"\[\[([^\]|]+)(?:\|([^\]]+))?\]\]")

# Plain symbols for the band and state emoji, in the Word/PDF only — the .md keeps the emoji.
EMOJI = {"🔥": "●●●", "⚠️": "●●○", "⚠": "●●○", "📎": "●○○", "🔴": "▼", "🟡": "◆", "🟢": "▲", "⚪": "○"}


def swap_emoji(text):
    for emoji, symbol in EMOJI.items():
        text = text.replace(emoji, symbol)
    return text


# Inside a ::: {custom-style="…"} box, pandoc styles paragraphs but not list items, so a list in a box
# would lose the box's shading. This filter turns such lists into paragraphs that start with "•" or the
# item number.
BOX_LISTS_FILTER = r"""
local function flatten(blocks)
  local out = {}
  for _, b in ipairs(blocks) do
    if b.t == 'BulletList' or b.t == 'OrderedList' then
      local n = 1
      if b.t == 'OrderedList' then n = (b.listAttributes and b.listAttributes.start) or b.start or 1 end
      for _, item in ipairs(b.content) do
        local marker = b.t == 'BulletList' and '•' or (tostring(n) .. '.')
        local first = true
        for _, ib in ipairs(item) do
          if ib.t == 'Plain' or ib.t == 'Para' then
            local inlines = ib.content
            if first then
              table.insert(inlines, 1, pandoc.Space())
              table.insert(inlines, 1, pandoc.Str(marker))
              first = false
            end
            table.insert(out, pandoc.Para(inlines))
          else
            table.insert(out, ib)
          end
        end
        n = n + 1
      end
    else
      table.insert(out, b)
    end
  end
  return out
end

function Div(el)
  if el.attributes['custom-style'] then
    el.content = flatten(el.content)
    return el
  end
end
"""


def read_frontmatter(text):
    """Return the flat key: value pairs of a YAML frontmatter block (enough for our keys)."""
    match = re.match(r"^---\r?\n(.*?)\r?\n---\r?\n", text, re.S)
    if not match:
        return {}
    meta = {}
    for line in match.group(1).splitlines():
        pair = re.match(r"^([A-Za-z_-]+):\s*(.*)$", line)
        if pair and pair.group(2):
            meta[pair.group(1)] = pair.group(2).strip().strip("\"'")
    return meta


def strip_wikilinks(text):
    return WIKILINK.sub(lambda m: m.group(2) or m.group(1).split("/")[-1], text)


def highlight_flag(pandoc):
    """pandoc 3.8 renamed --highlight-style to --syntax-highlighting."""
    out = subprocess.run([pandoc, "--version"], capture_output=True, text=True).stdout
    match = re.search(r"pandoc(?:\.exe)?\s+(\d+)\.(\d+)", out)
    version = (int(match.group(1)), int(match.group(2))) if match else (0, 0)
    return "--syntax-highlighting=tango" if version >= (3, 8) else "--highlight-style=tango"


MERMAID = re.compile(r"^```mermaid[ \t]*(.*?)[ \t]*\r?\n(.*?)^```[ \t]*$", re.S | re.M)


def diagram_width(png):
    """Width, in % of the text block, for a diagram. A narrow diagram is enlarged (text up to about 12 pt) to
    fill most of the page width instead of sitting in a thin column; a wide one is kept at 9.5 pt or more, even
    if that makes it wider than the page's comfortable share (at most 100%); and none is taller than ~420 pt."""
    with open(png, "rb") as handle:
        header = handle.read(24)
    width, height = struct.unpack(">II", header[16:24])  # rendered at scale 2
    native_w, native_h = width / 2, height / 2
    fill = 0.9 * 453 / native_w  # scale that would fill 90% of the text block (about 453 pt, 16 cm)
    scale = min(12 / 16, max(9.5 / 16, fill))  # pt per native pixel for a 16 px font
    width_pt, height_pt = native_w * scale, native_h * scale
    if height_pt > 420:
        width_pt *= 420 / height_pt
    return int(max(20, min(100, width_pt / 453 * 100)))


def render_mermaid(text, tmp):
    """Replace every ```mermaid block with an image rendered by mermaid-cli (mmdc)."""
    if not MERMAID.search(text):
        return text
    mmdc = shutil.which("mmdc")
    if not mmdc:
        sys.stderr.write("the source has mermaid diagrams but mmdc (mermaid-cli) is not installed\n")
        sys.exit(5)
    counter = [0]
    config = Path(tmp) / "mermaid-config.json"
    # natural size, not shrunk to the viewport: the figure's own width then tells how legible it is
    config.write_text(
        json.dumps({"flowchart": {"useMaxWidth": False}, "sequence": {"useMaxWidth": False}}), encoding="utf-8"
    )

    def one(match):
        counter[0] += 1
        mmd = Path(tmp) / f"diagram-{counter[0]}.mmd"
        png = Path(tmp) / f"diagram-{counter[0]}.png"
        mmd.write_text(match.group(2), encoding="utf-8")
        result = subprocess.run(
            [mmdc, "-i", str(mmd), "-o", str(png), "-c", str(config), "-s", "2", "-b", "white"],
            capture_output=True,
            text=True,
        )
        if result.returncode != 0 or not png.is_file():
            sys.stderr.write(f"mmdc failed on diagram {counter[0]}:\n{result.stderr or result.stdout}\n")
            sys.exit(6)
        caption = match.group(1).replace("]", ")").replace("[", "(")
        return f"![{caption}]({png.as_posix()}){{width={diagram_width(png)}%}}"

    return MERMAID.sub(one, text)


def run_pandoc(pandoc, source, docx_path, reference, toc):
    text = swap_emoji(strip_wikilinks(source.read_text(encoding="utf-8")))
    with tempfile.TemporaryDirectory() as tmp:
        text = render_mermaid(text, tmp)
        prepared = Path(tmp) / source.name
        prepared.write_text(text, encoding="utf-8")
        box_filter = Path(tmp) / "box-lists.lua"
        box_filter.write_text(BOX_LISTS_FILTER, encoding="utf-8")
        cmd = [
            pandoc,
            str(prepared),
            "--from=markdown+tex_math_dollars+pipe_tables+fenced_divs+bracketed_spans",
            "--to=docx",
            f"--reference-doc={reference}",
            f"--resource-path={source.parent}",
            f"--lua-filter={box_filter}",
            highlight_flag(pandoc),
            "-o",
            str(docx_path),
        ]
        if toc:
            cmd[2:2] = ["--toc", "--toc-depth=2"]
        result = subprocess.run(cmd, capture_output=True, text=True)
    if result.returncode != 0:
        sys.stderr.write(result.stderr)
        sys.exit(4)
    if result.stderr.strip():
        print(result.stderr.strip())


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def add_footer(docx_path, meta, fallback_title):
    """Footer on every page: '<title> — v<version> — <updated>' on the left, page number on the right."""
    parts = [meta.get("title", fallback_title)]
    if meta.get("version"):
        parts.append(f"v{meta['version']}")
    if meta.get("updated"):
        parts.append(meta["updated"])
    label = " — ".join(parts)

    doc = Document(str(docx_path))
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.text = label
        width = section.page_width - section.left_margin - section.right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        paragraph.add_run("\t")
        add_field(paragraph, "PAGE")
    doc.save(str(docx_path))


BOX_STYLES = ("Exam box", "Warning box", "Example box", "Note box")
FIGURE_CAPTION_STYLES = ("Image Caption", "Figure Caption", "Captioned Figure Caption", "Caption")
# A table is kept whole on one page only when it is clearly small: few rows and little text. Word cannot
# honour keep-with-next across a table taller than a page and would just push it down, wasting space.
MAX_KEPT_TABLE_ROWS = 12
MAX_KEPT_TABLE_CHARS = 1200


def set_keep_with_next(paragraph):
    """Keep this paragraph on the same page as the next one."""
    paragraph.paragraph_format.keep_with_next = True


def add_row_flag(row, tag):
    """Add an empty row property (w:cantSplit, w:tblHeader…) to a table row, once."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn(tag)) is None:
        tr_pr.append(OxmlElement(tag))


TABLE_BORDER = "8FA3BF"   # thin grid lines between every row and column
TABLE_HEADER_FILL = "1F3A5F"  # header row: dark blue, white bold text
TABLE_ROW_FILLS = ("EAF0F8", "FFFFFF")  # body rows alternate: light blue, white


def style_table(table):
    """Full grid, a coloured header row and body rows in alternating colours."""
    tbl_pr = table._tbl.tblPr
    old = tbl_pr.find(qn("w:tblBorders"))
    if old is not None:
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), TABLE_BORDER)
        borders.append(element)
    tbl_pr.append(borders)
    for index, row in enumerate(table.rows):
        fill = TABLE_HEADER_FILL if index == 0 else TABLE_ROW_FILLS[(index - 1) % 2]
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            old = tc_pr.find(qn("w:shd"))
            if old is not None:
                tc_pr.remove(old)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:val"), "clear")
            shade.set(qn("w:color"), "auto")
            shade.set(qn("w:fill"), fill)
            tc_pr.append(shade)
            if index == 0:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def tidy_layout(docx_path, section_breaks):
    """Keep tables and boxes from being cut, and start each top-level section on a new page."""
    doc = Document(str(docx_path))

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        style_table(table)
        rows = table.rows
        for row in rows:
            add_row_flag(row, "w:cantSplit")
        if rows:
            add_row_flag(rows[0], "w:tblHeader")
        chars = sum(len(cell.text) for row in rows for cell in row.cells)
        if len(rows) <= MAX_KEPT_TABLE_ROWS and chars <= MAX_KEPT_TABLE_CHARS:
            for row in rows[:-1]:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        set_keep_with_next(paragraph)

    # Figures (images, charts, diagrams) and their captions are centred, and a figure stays with its caption.
    for paragraph in doc.paragraphs:
        style = paragraph.style.name if paragraph.style is not None else ""
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_keep_with_next(paragraph)
        elif style in FIGURE_CAPTION_STYLES:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # A box made of several paragraphs stays in one piece.
    paragraphs = doc.paragraphs
    for current, following in zip(paragraphs, paragraphs[1:]):
        name = current.style.name if current.style is not None else ""
        if name in BOX_STYLES and following.style is not None and following.style.name == name:
            set_keep_with_next(current)

    if not section_breaks:
        for paragraph in paragraphs:
            if paragraph.style is not None and paragraph.style.name == "Heading 1":
                paragraph.paragraph_format.page_break_before = False

    doc.save(str(docx_path))


WORD_SCRIPT = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
    $doc = $word.Documents.Open($args[0], $false, $false)
    foreach ($toc in $doc.TablesOfContents) { $toc.Update() }
    $doc.Fields.Update() | Out-Null
    $doc.Save()
    $doc.SaveAs2($args[1], 17)
    $doc.Close($false)
} finally {
    $word.Quit()
}
"""


def make_pdf(docx_path, pdf_path):
    """Return None on success, or the reason the PDF was skipped."""
    if platform.system() == "Windows":
        shell = shutil.which("powershell") or shutil.which("pwsh")
        if shell:
            with tempfile.NamedTemporaryFile("w", suffix=".ps1", delete=False, encoding="utf-8") as script:
                script.write(WORD_SCRIPT)
            try:
                result = subprocess.run(
                    [shell, "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", script.name,
                     str(docx_path.resolve()), str(pdf_path.resolve())],
                    capture_output=True, text=True, timeout=300,
                )
            finally:
                os.unlink(script.name)
            if result.returncode == 0 and pdf_path.exists():
                return None
            word_error = (result.stderr or result.stdout).strip().splitlines()
            word_error = word_error[-1] if word_error else "unknown error"
        else:
            word_error = "PowerShell not found"
    else:
        word_error = "Microsoft Word automation is only available on Windows"

    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        result = subprocess.run(
            [office, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(docx_path)],
            capture_output=True, text=True, timeout=300,
        )
        if result.returncode == 0 and pdf_path.exists():
            return None
        return f"LibreOffice failed: {result.stderr.strip()}"
    return f"no PDF converter ({word_error}; LibreOffice not found)"


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", type=Path)
    parser.add_argument("--no-pdf", action="store_true")
    parser.add_argument("--reference", type=Path, help="reference.docx to use instead of the plugin's own")
    args = parser.parse_args()

    source = args.source.resolve()
    if not source.is_file() or source.suffix.lower() != ".md":
        sys.stderr.write(f"not a markdown file: {source}\n")
        sys.exit(2)

    pandoc = shutil.which("pandoc")
    if not pandoc:
        sys.stderr.write("pandoc not found on PATH\n")
        sys.exit(3)

    reference = args.reference or make_reference.cached_reference(pandoc)
    docx_path = source.with_suffix(".docx")
    pdf_path = source.with_suffix(".pdf")
    meta = read_frontmatter(source.read_text(encoding="utf-8"))

    toc = meta.get("toc", "true").lower() not in ("false", "no", "0")
    run_pandoc(pandoc, source, docx_path, reference, toc)
    add_footer(docx_path, meta, source.stem)
    tidy_layout(docx_path, meta.get("section-breaks", "true").lower() not in ("false", "no", "0"))
    print(f"DOCX: {docx_path}")

    if args.no_pdf:
        print("PDF: skipped (--no-pdf)")
        return
    reason = make_pdf(docx_path, pdf_path)
    print(f"PDF: {pdf_path}" if reason is None else f"PDF: skipped ({reason})")


if __name__ == "__main__":
    main()
